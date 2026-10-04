import docx
import re
import os

tex_file = "manuscript/manuscript.tex"
docx_file = "manuscript/manuscript.docx"

doc = docx.Document()

with open(tex_file, 'r', encoding='utf-8') as f:
    lines = f.readlines()

title = ""
author = ""
abstract = ""
in_abstract = False

for line in lines:
    if line.startswith(r"\title{"):
        title = line.replace(r"\title{", "").replace("}", "").strip()
        doc.add_heading(title, 0)
    elif line.startswith(r"\author{"):
        author = line.replace(r"\author{", "").replace("}", "").strip()
        doc.add_heading(author, 1)
    elif line.startswith(r"\begin{abstract}"):
        in_abstract = True
        doc.add_heading('Abstract', 1)
    elif line.startswith(r"\end{abstract}"):
        in_abstract = False
    elif in_abstract:
        if line.strip():
            doc.add_paragraph(line.strip())
    elif line.startswith(r"\section{"):
        sec = line.replace(r"\section{", "").replace("}", "").strip()
        doc.add_heading(sec, 1)
    elif line.startswith(r"\subsection{"):
        sec = line.replace(r"\subsection{", "").replace("}", "").strip()
        doc.add_heading(sec, 2)
    elif line.strip() and not line.strip().startswith("\\") and not line.strip().startswith("%"):
        # Very basic clean up of citations
        text = re.sub(r"\\citep\{.*?\}", "[Citation]", line.strip())
        text = re.sub(r"\\textit\{(.*?)\}", r"\1", text)
        doc.add_paragraph(text)

doc.save(docx_file)
print("Docx created successfully.")
