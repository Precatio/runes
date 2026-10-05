"""Akademisk rapport ur mätkorpusen.

Dokumentet byggs som en lista av block (rubrik, stycke, tabell, figur, lista) som sedan renderas till
Markdown, HTML, LaTeX och Word. Alla siffror, tabeller och figurer räknas fram här; en språkmodell får
som mest skriva inledning och diskussion utifrån de framräknade fakta (se api/routers/reports.py).
"""
from __future__ import annotations

import base64
import datetime
import html
import io
from collections import defaultdict

import numpy as np

from src.stats import METRICS, METRIC_LABELS, attribute, summarize, ward_clustering

DIGITS = {"apex_vinkel_deg": 1, "asymmetri_deg": 1, "spårdjup_mm": 2, "spårbredd_mm": 2,
          "djup_bredd_kvot": 2, "bottenradie_mm": 2, "ytråhet_mm": 3}
FEATURE_NAMES = {"rune": "runor", "ornament": "ornamentik", "unknown": "ej angivet"}

REFERENCES = [
    "Axelson, J. 1993. Mellansvenska runristare. Förteckning över signerade och attribuerade inskrifter. "
    "(Runrön 5.) Uppsala.",
    "Gräslund, A.-S. 1998. Ornamentiken som dateringsgrund för Upplands runstenar. I: Innskrifter og "
    "datering / Dating inscriptions. Trondheim, s. 73–91.",
    "Kitzler Åhfeldt, L. 2002. Work and Worship. Laser Scanner Analysis of Viking Age Rune Stones. "
    "Stockholms universitet.",
    "Kitzler Åhfeldt, L. & Imer, L. M. 2019. Rune Carvers and Sponsor Families on Bornholm. "
    "Danish Journal of Archaeology 8. https://doi.org/10.7146/dja.v8i0.113226",
    "Samnordisk runtextdatabas. Institutionen för nordiska språk, Uppsala universitet. "
    "http://www.nordiska.uu.se/forskn/samnord.htm",
]


# ---- block helpers -------------------------------------------------------------------------

def h(level: int, text: str) -> dict:
    return {"type": "heading", "level": level, "text": text}


def p(text: str, ai: bool = False) -> dict:
    return {"type": "paragraph", "text": text, "ai": ai}


def table(headers: list[str], rows: list[list[str]], caption: str) -> dict:
    return {"type": "table", "headers": headers, "rows": rows, "caption": caption}


def figure(png_b64: str, caption: str, name: str) -> dict:
    return {"type": "figure", "png": png_b64, "caption": caption, "name": name}


def bullets(items: list[str]) -> dict:
    return {"type": "list", "items": items}


def fmt(value, digits=2) -> str:
    if value is None or (isinstance(value, float) and not np.isfinite(value)):
        return "–"
    return f"{value:.{digits}f}".replace(".", ",")


# ---- analysis ------------------------------------------------------------------------------

def _stone_means(entry: dict) -> list[float] | None:
    m = entry.get("means") or {}
    if any(m.get(k) is None for k in METRICS):
        return None
    return [float(m[k]) for k in METRICS]


def build_facts(entries: list[dict], rundata_lookup, scope: dict) -> dict:
    """Allt som rapporten påstår, framräknat ur korpusposterna och Rundata."""
    stones = []
    for e in entries:
        rec = rundata_lookup(e.get("signum") or "")
        certain = [c["name"] for c in (rec or {}).get("carvers", []) if c["kind"] in ("S", "A") and not c["uncertain"]]
        stones.append({
            "entry": e,
            "signum": rec["signum"] if rec else e.get("signum"),
            "place": (rec or {}).get("place", ""),
            "style": (rec or {}).get("style"),
            "dating": (rec or {}).get("dating", ""),
            "carver": certain[0] if len(certain) == 1 else None,
            "carver_text": ", ".join(f"{c['name']} ({c['kind']})" for c in (rec or {}).get("carvers", [])
                                     if c["kind"] in ("S", "A")) or "–",
            "means": _stone_means(e),
        })

    groups: dict[str, list[dict]] = defaultdict(list)
    for s in stones:
        if s["carver"] and s["means"]:
            groups[s["carver"]].append(s)

    carver_rows = []
    for name, ss in sorted(groups.items(), key=lambda kv: -len(kv[1])):
        stats = {m: summarize([x["means"][i] for x in ss]) for i, m in enumerate(METRICS)}
        carver_rows.append({"carver": name, "n": len(ss), "stats": stats})

    reference = [{"group": s["carver"], "label": s["signum"], "values": s["means"]}
                 for s in stones if s["carver"] and s["means"]]
    evaluation = None
    if len({r["group"] for r in reference}) >= 2:
        evaluation = attribute(reference, reference[0]["values"], min_per_group=2).get("evaluation")

    usable = [s for s in stones if s["means"]]
    clustering = None
    if len(usable) >= 3:
        labels = [f"{s['signum']}" + (f" ({s['carver']})" if s["carver"] else "") for s in usable]
        clustering = {"labels": labels, **ward_clustering(labels, np.array([s["means"] for s in usable]))}

    versions = sorted({(s["entry"].get("method_version") or "okänd") for s in stones})
    contributors = sorted({f"{s['entry'].get('contributorName', '')}"
                           + (f", {s['entry'].get('institution')}" if s["entry"].get("institution") else "")
                           for s in stones if s["entry"].get("contributorName")})
    n_slices = sum(len(s["entry"].get("slices") or []) for s in stones)
    feature_types = sorted({s["entry"].get("feature_type", "unknown") for s in stones})
    return {
        "scope": scope, "stones": stones, "carver_rows": carver_rows, "evaluation": evaluation,
        "clustering": clustering, "method_versions": versions, "contributors": contributors,
        "n_slices": n_slices, "feature_types": feature_types,
        "date": datetime.date.today().isoformat(),
    }


def facts_text(facts: dict) -> str:
    """Kortfattad sammanställning för språkmodellen. Bara detta får användas i AI-text."""
    lines = [f"Urval: {facts['scope'].get('title') or facts['scope'].get('type')}",
             f"Antal stenar: {len(facts['stones'])}; antal tvärsnitt: {facts['n_slices']}; "
             f"spårtyper: {', '.join(FEATURE_NAMES.get(f, f) for f in facts['feature_types'])}",
             f"Mätmetod(er): {', '.join(facts['method_versions'])}"]
    for s in facts["stones"]:
        m = s["means"]
        lines.append(f"- {s['signum']} ({s['place']}), ristare enligt Rundata: {s['carver_text']}, stil {s['style'] or '–'}, "
                     f"datering {s['dating'] or '–'}" + (f", V-vinkel {m[0]:.1f}°, djup {m[2]:.2f} mm, bredd {m[3]:.2f} mm" if m else ""))
    for r in facts["carver_rows"]:
        a = r["stats"]["apex_vinkel_deg"]
        lines.append(f"Ristare {r['carver']}: {r['n']} stenar, V-vinkel medel {fmt(a['mean'], 1)}° (SD {fmt(a['sd'], 1)})")
    ev = facts["evaluation"]
    if ev:
        lines.append(f"Korsvaliderad attribuering: rätt ristare först i {ev['top1_accuracy']:.0%} av {ev['n_stones']} stenar "
                     f"({ev['n_groups']} ristare, slumpnivå {ev['chance_top1']:.0%}).")
    else:
        lines.append("Attribueringen kunde inte utvärderas (för få ristare med minst två uppmätta stenar).")
    return "\n".join(lines)


# ---- figures -------------------------------------------------------------------------------

def _png(fig) -> str:
    buf = io.BytesIO()
    fig.savefig(buf, format="png", dpi=200, bbox_inches="tight")
    import matplotlib.pyplot as plt
    plt.close(fig)
    return base64.b64encode(buf.getvalue()).decode("ascii")


def angle_figure(facts: dict) -> str | None:
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    stones = [s for s in facts["stones"] if s["means"]]
    if len(stones) < 2:
        return None
    fig, ax = plt.subplots(figsize=(7, 0.35 * len(stones) + 1.2))
    order = sorted(stones, key=lambda s: (s["carver"] or "ö", s["means"][0]))
    for i, s in enumerate(order):
        summ = (s["entry"].get("summary") or {}).get("apex_vinkel_deg") or {}
        mean, sd = s["means"][0], summ.get("sd") or 0
        ax.errorbar(mean, i, xerr=sd, fmt="o", color="#b7410e" if s["carver"] else "#64748b", capsize=3)
    ax.set_yticks(range(len(order)))
    ax.set_yticklabels([s["signum"] + (f" – {s['carver']}" if s["carver"] else "") for s in order], fontsize=8)
    ax.set_xlabel("V-vinkel (°), medel ± SD över tvärsnitt")
    ax.grid(axis="x", alpha=0.3)
    return _png(fig)


def dendrogram_figure(facts: dict) -> str | None:
    c = facts["clustering"]
    if not c or "icoord" not in c:
        return None
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    fig, ax = plt.subplots(figsize=(7, 3.5))
    for xs, ys in zip(c["icoord"], c["dcoord"]):
        ax.plot(xs, ys, color="#0f172a", lw=1)
    ax.set_xticks([5 + 10 * i for i in range(len(c["order"]))])
    ax.set_xticklabels(c["order"], rotation=60, ha="right", fontsize=8)
    ax.set_ylabel("Avstånd (Ward)")
    ax.spines[["top", "right"]].set_visible(False)
    return _png(fig)


# ---- document ------------------------------------------------------------------------------

def build_document(facts: dict, author: str, institution: str, ai: dict | None) -> list[dict]:
    scope = facts["scope"]
    title = scope.get("title") or "Huggteknik i runstenskorpusen"
    stones = facts["stones"]
    ai = ai or {}
    blocks = [
        {"type": "title", "text": title, "author": author, "institution": institution, "date": facts["date"]},
        h(1, "Sammanfattning"),
        p(ai.get("abstract") or
          f"Rapporten sammanställer huggspårsmätningar från {len(stones)} runstenar ({facts['n_slices']} tvärsnitt) "
          f"i Bifrosts mätkorpus, med ristaruppgifter ur Samnordisk runtextdatabas.", ai=bool(ai.get("abstract"))),
        h(1, "1. Inledning"),
        p(ai.get("introduction") or
          "Syftet är att beskriva huggtekniken i det valda urvalet och pröva i vilken mån uppmätta spårmått "
          "skiljer ristare åt. Frågeställningarna formuleras av författaren.", ai=bool(ai.get("introduction"))),
        h(1, "2. Material"),
        p(f"Materialet består av {len(stones)} uppmätta stenar. Ristaruppgifter, datering och stilgrupp är hämtade "
          "ur Samnordisk runtextdatabas; S betecknar signerad och A attribuerad inskrift. Mätningarna är publicerade "
          "i mätkorpusen under CC BY 4.0."),
        table(["Signum", "Plats", "Ristare (Rundata)", "Stil", "Datering", "Spår", "Snitt", "Bidragsgivare"],
              [[s["signum"], s["place"], s["carver_text"], s["style"] or "–", s["dating"] or "–",
                FEATURE_NAMES.get(s["entry"].get("feature_type", ""), "–"), str(len(s["entry"].get("slices") or [])),
                s["entry"].get("contributorName", "–")] for s in stones],
              "Tabell 1. Uppmätta stenar."),
        h(1, "3. Metod"),
        p("Tvärsnitt genom spåren har tagits ur 3D-skanningar med Bifrost "
          f"(mätmetod {', '.join(facts['method_versions'])}). I varje tvärsnitt anpassas spårväggarna med linjär "
          "regression mellan 20 och 80 procent av spårdjupet; V-vinkeln är öppningsvinkeln mellan väggarnas linjer, "
          "bredden mäts där linjerna når stenytan och djupet från spårkanten till botten. Metoden har validerats "
          "på syntetiska spår med kända mått (vinkelfel inom ±0,6°). Varje sten representeras av medelvärdet över "
          "sina tvärsnitt."),
        p("Stenarna jämförs med hierarkisk klustring (Wards metod, euklidiska avstånd på standardiserade "
          "medelvärden), i likhet med Kitzler Åhfeldts analyser av huggteknik. Ristarattribuering prövas med "
          "Mahalanobisavstånd till ristarnas medelvärden och utvärderas med lämna-en-ute-korsvalidering. "
          "Bara stenar med en säker signerad eller attribuerad ristare i Rundata används som referens."),
    ]
    if len(facts["method_versions"]) > 1:
        blocks.append(p("Obs: materialet innehåller mätningar gjorda med olika metodversioner, vilket begränsar "
                        "jämförbarheten. Mätningar med råprofiler kan räknas om med den aktuella metoden."))

    blocks.append(h(1, "4. Resultat"))
    metric_headers = ["Ristare", "n"] + [METRIC_LABELS[m] for m in METRICS[:4]]
    if facts["carver_rows"]:
        blocks.append(table(metric_headers,
                            [[r["carver"], str(r["n"])] + [f"{fmt(r['stats'][m]['mean'], DIGITS[m])} ± {fmt(r['stats'][m]['sd'], DIGITS[m])}"
                                                            for m in METRICS[:4]] for r in facts["carver_rows"]],
                            "Tabell 2. Medelvärde ± SD över stenar per ristare (endast säkra attribueringar)."))
    else:
        blocks.append(p("Inga stenar i urvalet har en säker ristare i Rundata, så inga ristarprofiler redovisas."))
    blocks.append(table(["Signum"] + [METRIC_LABELS[m] for m in METRICS],
                        [[s["signum"]] + [fmt(v, DIGITS[m]) for v, m in zip(s["means"], METRICS)]
                         for s in stones if s["means"]],
                        "Tabell 3. Medelvärden per sten."))
    fig1 = angle_figure(facts)
    if fig1:
        blocks.append(figure(fig1, "Figur 1. V-vinkel per sten (medel ± SD över tvärsnitt).", "figur1_v-vinkel.png"))
    fig2 = dendrogram_figure(facts)
    if fig2:
        blocks.append(figure(fig2, "Figur 2. Hierarkisk klustring av stenarna (Ward).", "figur2_klustring.png"))
    ev = facts["evaluation"]
    if ev:
        blocks.append(p(f"Med korsvalidering hamnade rätt ristare först för {ev['top1_accuracy'] * 100:.0f} procent av "
                        f"{ev['n_stones']} stenar och bland de tre första för {ev['top3_accuracy'] * 100:.0f} procent "
                        f"({ev['n_groups']} ristare; slumpnivå {ev['chance_top1'] * 100:.0f} procent)."))
    else:
        blocks.append(p("Attribueringen kunde inte utvärderas, eftersom urvalet har färre än två ristare med minst "
                        "två uppmätta stenar."))

    blocks.append(h(1, "5. Diskussion"))
    blocks.append(p(ai.get("discussion") or
                    "Diskussionen skrivs av författaren utifrån resultaten ovan.", ai=bool(ai.get("discussion"))))
    blocks.append(h(2, "Begränsningar"))
    blocks.append(bullets([
        "Spårmåtten påverkas av vittring, bergart, skanningens upplösning och var tvärsnitten läggs.",
        "Runor och ornamentik huggs ofta olika och bör jämföras var för sig.",
        "Ristaruppgifterna i Rundata är hypoteser i litteraturen och inte facit.",
        "Få stenar per ristare ger osäkra profiler; korsvalideringens siffror bör läsas med det i åtanke.",
    ]))
    blocks.append(h(1, "Referenser"))
    blocks.append(bullets(REFERENCES + [
        "Bifrost, version 2.0. Programvara. https://github.com/Precatio/runes",
    ]))
    if facts["contributors"]:
        blocks.append(h(2, "Bidragsgivare till mätdata"))
        blocks.append(bullets(facts["contributors"]))
    if any(b.get("ai") for b in blocks):
        blocks.append(p("Avsnitt markerade som AI-genererade är formulerade av en språkmodell utifrån de framräknade "
                        "resultaten och ska granskas av författaren."))
    return blocks


# ---- renderers -----------------------------------------------------------------------------

def to_markdown(blocks: list[dict]) -> str:
    out = []
    for b in blocks:
        t = b["type"]
        if t == "title":
            out += [f"# {b['text']}", "", f"{b['author']}" + (f", {b['institution']}" if b["institution"] else ""),
                    f"{b['date']}", ""]
        elif t == "heading":
            out += ["#" * (b["level"] + 1) + " " + b["text"], ""]
        elif t == "paragraph":
            out += [(b["text"] + (" *[AI-genererat]*" if b.get("ai") else "")), ""]
        elif t == "list":
            out += [f"- {i}" for i in b["items"]] + [""]
        elif t == "table":
            out += ["| " + " | ".join(b["headers"]) + " |", "|" + "---|" * len(b["headers"])]
            out += ["| " + " | ".join(c.replace("|", "/") for c in r) + " |" for r in b["rows"]]
            out += ["", f"*{b['caption']}*", ""]
        elif t == "figure":
            out += [f"![{b['caption']}]({b['name']})", "", f"*{b['caption']}*", ""]
        elif t == "inscription":
            out += [f"> **{b['transliteration']}**", ">"]
            if b["normalization"]:
                out += [f"> *{b['normalization']}*", ">"]
            if b["normalization_ows"]:
                out += [f"> *{b['normalization_ows']}*", ">"]
            if b["translation"]:
                q = ("“", "”") if b.get("lang") == "en" else ("”", "”")
                out += [f"> {q[0]}{b['translation']}{q[1]}"]
            out += [""]
    return "\n".join(out)


def to_html(blocks: list[dict]) -> str:
    e = html.escape
    out = []
    for b in blocks:
        t = b["type"]
        if t == "title":
            out.append(f"<h1>{e(b['text'])}</h1><p><em>{e(b['author'])}{', ' + e(b['institution']) if b['institution'] else ''} · {e(b['date'])}</em></p>")
        elif t == "heading":
            out.append(f"<h{b['level'] + 1}>{e(b['text'])}</h{b['level'] + 1}>")
        elif t == "paragraph":
            tag = ' <span class="ai-badge">AI-genererat</span>' if b.get("ai") else ""
            out.append(f"<p>{e(b['text'])}{tag}</p>")
        elif t == "list":
            out.append("<ul>" + "".join(f"<li>{e(i)}</li>" for i in b["items"]) + "</ul>")
        elif t == "table":
            head = "".join(f"<th>{e(x)}</th>" for x in b["headers"])
            rows = "".join("<tr>" + "".join(f"<td>{e(c)}</td>" for c in r) + "</tr>" for r in b["rows"])
            out.append(f"<table><thead><tr>{head}</tr></thead><tbody>{rows}</tbody></table><p class=\"caption\">{e(b['caption'])}</p>")
        elif t == "figure":
            out.append(f"<figure><img src=\"data:image/png;base64,{b['png']}\" alt=\"{e(b['caption'])}\"/>"
                       f"<figcaption>{e(b['caption'])}</figcaption></figure>")
        elif t == "inscription":
            parts = [f"<p><strong>{e(b['transliteration'])}</strong></p>"]
            parts += [f"<p><em>{e(x)}</em></p>" for x in (b["normalization"], b["normalization_ows"]) if x]
            if b["translation"]:
                q = ("“", "”") if b.get("lang") == "en" else ("”", "”")
                parts.append(f"<p>{q[0]}{e(b['translation'])}{q[1]}</p>")
            out.append(f"<blockquote class=\"inscription\">{''.join(parts)}</blockquote>")
    return "\n".join(out)


def _tex(s: str) -> str:
    rep = {"\\": r"\textbackslash{}", "&": r"\&", "%": r"\%", "$": r"\$", "#": r"\#", "_": r"\_",
           "{": r"\{", "}": r"\}", "~": r"\textasciitilde{}", "^": r"\textasciicircum{}"}
    return "".join(rep.get(ch, ch) for ch in s)


def to_latex(blocks: list[dict]) -> str:
    out = [r"\documentclass[11pt]{article}", r"\usepackage[utf8]{inputenc}", r"\usepackage[T1]{fontenc}",
           r"\usepackage[swedish]{babel}", r"\usepackage{graphicx}", r"\usepackage{booktabs}",
           r"\usepackage{longtable}", r"\begin{document}"]
    for b in blocks:
        t = b["type"]
        if t == "title":
            out += [rf"\title{{{_tex(b['text'])}}}",
                    rf"\author{{{_tex(b['author'])}\\{_tex(b['institution'])}}}", rf"\date{{{b['date']}}}", r"\maketitle"]
        elif t == "heading":
            cmd = {1: "section*", 2: "subsection*"}.get(b["level"], "subsubsection*")
            out.append(rf"\{cmd}{{{_tex(b['text'])}}}")
        elif t == "paragraph":
            out += [_tex(b["text"]) + (r" \emph{[AI-genererat]}" if b.get("ai") else ""), ""]
        elif t == "list":
            out += [r"\begin{itemize}"] + [rf"\item {_tex(i)}" for i in b["items"]] + [r"\end{itemize}"]
        elif t == "table":
            cols = "l" * len(b["headers"])
            out += [rf"\begin{{longtable}}{{{cols}}}", r"\toprule", " & ".join(_tex(x) for x in b["headers"]) + r" \\",
                    r"\midrule"] + [" & ".join(_tex(c) for c in r) + r" \\" for r in b["rows"]] + \
                   [r"\bottomrule", rf"\caption{{{_tex(b['caption'])}}}", r"\end{longtable}"]
        elif t == "figure":
            out += [r"\begin{figure}[h]", r"\centering", rf"\includegraphics[width=\linewidth]{{{b['name']}}}",
                    rf"\caption{{{_tex(b['caption'])}}}", r"\end{figure}"]
        elif t == "inscription":
            out += [r"\begin{quote}", rf"\textbf{{{_tex(b['transliteration'])}}}\\[0.5ex]"]
            out += [rf"\emph{{{_tex(x)}}}\\[0.5ex]" for x in (b["normalization"], b["normalization_ows"]) if x]
            if b["translation"]:
                out.append(rf"''{_tex(b['translation'])}''")
            out.append(r"\end{quote}")
    out.append(r"\end{document}")
    return "\n".join(out)


def to_docx(blocks: list[dict]) -> bytes:
    from docx import Document
    from docx.shared import Inches

    doc = Document()
    for b in blocks:
        t = b["type"]
        if t == "title":
            doc.add_heading(b["text"], level=0)
            doc.add_paragraph(f"{b['author']}" + (f", {b['institution']}" if b["institution"] else "") + f"\n{b['date']}")
        elif t == "heading":
            doc.add_heading(b["text"], level=b["level"])
        elif t == "paragraph":
            par = doc.add_paragraph(b["text"])
            if b.get("ai"):
                par.add_run(" [AI-genererat]").italic = True
        elif t == "list":
            for i in b["items"]:
                doc.add_paragraph(i, style="List Bullet")
        elif t == "table":
            tbl = doc.add_table(rows=1, cols=len(b["headers"]))
            tbl.style = "Light Grid Accent 1"
            for cell, text in zip(tbl.rows[0].cells, b["headers"]):
                cell.text = text
            for r in b["rows"]:
                for cell, text in zip(tbl.add_row().cells, r):
                    cell.text = text
            doc.add_paragraph(b["caption"]).runs[0].italic = True
        elif t == "figure":
            doc.add_picture(io.BytesIO(base64.b64decode(b["png"])), width=Inches(6))
            doc.add_paragraph(b["caption"]).runs[0].italic = True
        elif t == "inscription":
            doc.add_paragraph(style="Quote").add_run(b["transliteration"]).bold = True
            for x in (b["normalization"], b["normalization_ows"]):
                if x:
                    doc.add_paragraph(style="Quote").add_run(x).italic = True
            if b["translation"]:
                q = ("“", "”") if b.get("lang") == "en" else ("”", "”")
                doc.add_paragraph(f"{q[0]}{b['translation']}{q[1]}", style="Quote")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
